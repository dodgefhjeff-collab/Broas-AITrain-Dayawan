# -*- coding: utf-8 -*-
"""把客户操作手册 Markdown 排成 A4 Word。"""
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor, Emu

ROOT = Path(__file__).resolve().parent
MD = ROOT / "客户操作手册.md"
OUT = ROOT / "客户操作手册-排版.docx"
NAVY = RGBColor(0x1A, 0x65, 0x8F)
IMG_WIDTH = Cm(15.5)


def set_run_font(run, size, bold=False, color=None):
    run.font.name = "微软雅黑"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def set_style_font(style, size, bold=False, color=None):
    style.font.name = "微软雅黑"
    style.font.size = Pt(size)
    style.font.bold = bold
    if color is not None:
        style.font.color.rgb = color
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), "微软雅黑")
    rfonts.set(qn("w:hAnsi"), "微软雅黑")
    rfonts.set(qn("w:eastAsia"), "微软雅黑")


def add_text_runs(paragraph, text, size=12, bold=False, color=None):
    parts = re.split(r"(\*\*.+?\*\*)", text)
    for part in parts:
        if not part:
            continue
        is_bold = part.startswith("**") and part.endswith("**")
        content = part[2:-2] if is_bold else part
        run = paragraph.add_run(content)
        set_run_font(run, size, bold or is_bold, color)


def shade_cell(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_border(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "BFBFBF")
        borders.append(el)
    tcPr.append(borders)


def add_page_number(paragraph):
    run = paragraph.add_run()
    set_run_font(run, 9, color=RGBColor(0x66, 0x66, 0x66))

    def fld(char_type):
        el = OxmlElement("w:fldChar")
        el.set(qn("w:fldCharType"), char_type)
        return el

    begin = fld("begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = fld("end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)


def setup_section(doc):
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.0)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)

    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = hp.add_run("室内训练机操作手册")
    set_run_font(run, 9, color=RGBColor(0x66, 0x66, 0x66))

    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_number(fp)


def add_picture(doc, rel):
    path = (ROOT / rel).resolve()
    if not path.exists():
        p = doc.add_paragraph()
        add_text_runs(p, f"（缺少图片：{rel}）", 10, color=RGBColor(0x99, 0x33, 0x33))
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    run = p.add_run()
    run.add_picture(str(path), width=IMG_WIDTH)


def add_table(doc, rows):
    if not rows:
        return
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.autofit = True
    table.allow_autofit = True
    tbl = table._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else OxmlElement("w:tblPr")
    tblW = OxmlElement("w:tblW")
    tblW.set(qn("w:w"), "5000")
    tblW.set(qn("w:type"), "pct")
    # replace existing tblW
    old = tblPr.find(qn("w:tblW"))
    if old is not None:
        tblPr.remove(old)
    tblPr.append(tblW)

    for i, row in enumerate(rows):
        for j in range(cols):
            cell = table.cell(i, j)
            text = row[j] if j < len(row) else ""
            cell.text = ""
            cp = cell.paragraphs[0]
            cp.paragraph_format.space_before = Pt(0)
            cp.paragraph_format.space_after = Pt(0)
            cp.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            add_text_runs(cp, text, 10.5, color=RGBColor(0xFF, 0xFF, 0xFF) if i == 0 else None)
            if i == 0:
                for run in cp.runs:
                    run.bold = True
                    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                shade_cell(cell, "1A658F")
            elif i % 2 == 0:
                shade_cell(cell, "F4F8FA")
            set_cell_border(cell)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def build():
    doc = Document()
    setup_section(doc)
    set_style_font(doc.styles["Normal"], 12)
    doc.styles["Normal"].paragraph_format.space_after = Pt(6)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.15
    for style_name, size in (("Heading 1", 16), ("Heading 2", 13)):
        st = doc.styles[style_name]
        set_style_font(st, size, True, NAVY)
        st.paragraph_format.space_before = Pt(14 if size == 16 else 10)
        st.paragraph_format.space_after = Pt(6)

    lines = MD.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if not stripped or stripped == "---":
            i += 1
            continue
        if stripped.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                raw = lines[i].strip().strip("|")
                cells = [c.strip() for c in raw.split("|")]
                if not all(re.fullmatch(r":?-{3,}:?", c.replace(" ", "")) or c.replace(" ", "") == "---" for c in cells):
                    if not re.match(r"^[\s|:-]+$", lines[i].strip()):
                        rows.append(cells)
                i += 1
            add_table(doc, rows)
            continue
        img = re.match(r"!\[(.*?)\]\((.*?)\)", stripped)
        if img:
            add_picture(doc, img.group(2))
            i += 1
            continue
        if stripped.startswith("## "):
            doc.add_heading(stripped[3:].strip(), level=1)
            i += 1
            continue
        if stripped.startswith("### "):
            doc.add_heading(stripped[4:].strip(), level=2)
            i += 1
            continue
        if stripped.startswith("# "):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(8)
            add_text_runs(p, stripped[2:].strip(), 22, True, NAVY)
            i += 1
            continue
        if stripped.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.4)
            add_text_runs(p, stripped[2:].strip(), 11, color=RGBColor(0x55, 0x55, 0x55))
            i += 1
            continue
        if stripped.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(2)
            add_text_runs(p, stripped[2:].strip(), 12)
            i += 1
            continue
        mnum = re.match(r"^(\d+)\.\s+(.*)", stripped)
        if mnum:
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.space_after = Pt(2)
            add_text_runs(p, mnum.group(2), 12)
            i += 1
            continue
        p = doc.add_paragraph()
        add_text_runs(p, stripped, 12)
        i += 1

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()

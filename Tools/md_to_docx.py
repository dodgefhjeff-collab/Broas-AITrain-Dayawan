import re
import sys
from typing import List

from docx import Document
from docx.shared import Pt


def _parse_table_row(line: str) -> List[str]:
    # Markdown table rows look like: | a | b | c |
    parts = [p.strip() for p in line.strip().strip("|").split("|")]
    return parts


def _is_table_separator(line: str) -> bool:
    # Separator example: |---|:---:|---:|
    cells = [c.strip() for c in line.strip().strip("|").split("|")]
    if not cells:
        return False
    for c in cells:
        if c == "":
            return False
        c2 = c.replace(":", "").replace("-", "")
        if c2 != "":
            return False
    return True


def _apply_code_style(run) -> None:
    run.font.name = "Consolas"
    run.font.size = Pt(9)


def markdown_to_docx(input_md_path: str, output_docx_path: str) -> None:
    document = Document()

    with open(input_md_path, "r", encoding="utf-8") as f:
        # Keep line breaks behavior close to original; we process by lines.
        lines = f.read().splitlines()

    i = 0
    code_mode = False
    while i < len(lines):
        raw = lines[i]
        line = raw.rstrip("\r\n")
        stripped = line.strip()

        # Code fences
        if stripped.startswith("```"):
            code_mode = not code_mode
            i += 1
            continue

        if code_mode:
            p = document.add_paragraph()
            run = p.add_run(line)
            _apply_code_style(run)
            i += 1
            continue

        # Empty line
        if stripped == "":
            document.add_paragraph("")
            i += 1
            continue

        # Headings
        if stripped.startswith("#"):
            m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
            if m:
                level = min(len(m.group(1)), 6)
                text = m.group(2).strip()
                p = document.add_paragraph(text)
                p.style = f"Heading {level}"
                i += 1
                continue

        # Blockquote
        if stripped.startswith(">"):
            text = stripped.lstrip(">").strip()
            document.add_paragraph(text)
            i += 1
            continue

        # Markdown table (very small subset)
        if stripped.startswith("|") and i + 1 < len(lines) and lines[i + 1].strip().startswith("|"):
            if _is_table_separator(lines[i + 1]):
                header = _parse_table_row(lines[i])
                i += 2
                rows: List[List[str]] = []
                while i < len(lines) and lines[i].strip().startswith("|"):
                    row = _parse_table_row(lines[i])
                    # Pad or truncate to match header columns
                    if len(row) < len(header):
                        row = row + [""] * (len(header) - len(row))
                    elif len(row) > len(header):
                        row = row[: len(header)]
                    rows.append(row)
                    i += 1

                # Create docx table
                table = document.add_table(rows=len(rows) + 1, cols=len(header))
                table.style = "Table Grid"
                for col_idx, h in enumerate(header):
                    table.cell(0, col_idx).text = h
                for r_idx, row in enumerate(rows, start=1):
                    for c_idx, val in enumerate(row):
                        table.cell(r_idx, c_idx).text = val
                continue

        # Bullet / numbered lists
        if re.match(r"^(-|\*)\s+.+$", stripped):
            text = re.sub(r"^(-|\*)\s+", "", stripped).strip()
            p = document.add_paragraph(text)
            p.style = "List Bullet"
            i += 1
            continue
        if re.match(r"^\d+\.\s+.+$", stripped):
            text = re.sub(r"^\d+\.\s+", "", stripped).strip()
            p = document.add_paragraph(text)
            p.style = "List Number"
            i += 1
            continue

        # Horizontal rules (optional)
        if re.fullmatch(r"-{3,}", stripped):
            document.add_paragraph("")
            i += 1
            continue

        # Default paragraph
        document.add_paragraph(stripped)
        i += 1

    document.save(output_docx_path)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print("Usage: python md_to_docx.py <input_md_path> <output_docx_path>")
        sys.exit(2)

    input_md_path = sys.argv[1]
    output_docx_path = sys.argv[2]
    markdown_to_docx(input_md_path, output_docx_path)

